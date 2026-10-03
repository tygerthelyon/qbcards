"""md2docx.py -- turn these notes into .docx in the style of my World Hist doc (bullets, bold key terms)."""
import re, sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_COLOR_INDEX

def runs(par, text):
    # **bold**, *italic*, ***both***
    for tok in re.split(r'(\*\*\*[^*]+\*\*\*|\*\*[^*]+\*\*|\*[^*]+\*)', text):
        if not tok: continue
        if tok.startswith('***'): r = par.add_run(tok[3:-3]); r.bold = r.italic = True
        elif tok.startswith('**'): r = par.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith('*'): r = par.add_run(tok[1:-1]); r.italic = True
        else: par.add_run(tok)

def convert(src, dst):
    doc = Document()
    st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(12)
    lines = open(src, encoding='utf-8').read().split('\n')
    i = 0; buf = None
    while i < len(lines):
        ln = lines[i]
        # join wrapped continuation lines of a bullet/paragraph
        while i + 1 < len(lines) and lines[i+1].startswith('  ') and not lines[i+1].lstrip().startswith(('- ', '|')) and ln.strip():
            i += 1; ln = ln.rstrip() + ' ' + lines[i].strip()
        s = ln.strip()
        if not s or s == '---': i += 1; continue
        if s.startswith('# '): h = doc.add_heading(s[2:], 0)
        elif s.startswith('## '):
            p = doc.add_paragraph(); r = p.add_run(s[3:]); r.bold = True; r.font.size = Pt(14)
            r.font.highlight_color = WD_COLOR_INDEX.TURQUOISE if False else WD_COLOR_INDEX.GRAY_25
        elif s.startswith('### '):
            p = doc.add_paragraph(); r = p.add_run(s[4:]); r.bold = True; r.underline = True
        elif s.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(set(c) <= set('-: ') for c in cells): rows.append(cells)
                i += 1
            t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = 'Light Grid Accent 1'
            for r_i, row in enumerate(rows):
                for c_i, c in enumerate(row): runs(t.cell(r_i, c_i).paragraphs[0], c)
            continue
        elif re.match(r'^\s*- ', ln):
            indent = (len(ln) - len(ln.lstrip())) // 2
            p = doc.add_paragraph(style='List Bullet 2' if indent else 'List Bullet'); runs(p, s[2:])
        else:
            p = doc.add_paragraph(); runs(p, s)
        i += 1
    doc.save(dst)

for f in sys.argv[1:]:
    convert(f, f.replace('.md', '.docx')); print('wrote', f.replace('.md', '.docx'))
